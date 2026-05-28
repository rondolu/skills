#!/usr/bin/env python3

import json
import os
import shutil

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PPTX_DIR = os.path.join(BASE_DIR, "skills", "pptx")
ASSET_SOURCE_DIR = os.path.join(PPTX_DIR, "assets", "backgrounds")
HTML_SOURCE = os.path.join(
    BASE_DIR,
    "skills",
    "frontend-design",
    "yahoo-fantasy-baseball-analytics.html",
)
YAML_SOURCE = os.path.join(
    PPTX_DIR,
    "fetch_template_with_backgrounds_editorial.yaml",
)
COVER_BG_SOURCE = os.path.join(ASSET_SOURCE_DIR, "cover-background.png")
CONTENT_BG_SOURCE = os.path.join(ASSET_SOURCE_DIR, "content-background.png")

DECK_NAME = "Yahoo-Fantasy-Baseball-Analyzer-Editorial-Deck"
BUNDLE_DIR = os.path.join(BASE_DIR, DECK_NAME)
BUNDLE_ASSET_DIR = os.path.join(BUNDLE_DIR, "assets")
BUNDLE_BG_DIR = os.path.join(BUNDLE_ASSET_DIR, "backgrounds")
OUTPUT_PPTX = os.path.join(BUNDLE_DIR, f"{DECK_NAME}.pptx")
COVER_BG = os.path.join(BUNDLE_BG_DIR, "cover-background.png")
CONTENT_BG = os.path.join(BUNDLE_BG_DIR, "content-background.png")

DISPLAY_FONT = "Arial"
MONO_FONT = "Courier New"

SLIDE_WIDTH = Inches(10)
SLIDE_HEIGHT = Inches(5.625)

SAFE_X = Inches(0.90)
SAFE_Y = Inches(0.45)
SAFE_W = Inches(8.60)
SAFE_H = Inches(4.60)

HEADER_X = SAFE_X
HEADER_W = SAFE_W
EYEBROW_Y = Inches(0.52)
TITLE_Y = Inches(0.86)
CONTENT_Y = Inches(1.50)
CONTENT_H = Inches(3.75)


def color(hex_value):
    red = int(hex_value[0:2], 16)
    green = int(hex_value[2:4], 16)
    blue = int(hex_value[4:6], 16)
    return RGBColor(red, green, blue)


COLORS = {
    "white": color("FFFFFF"),
    "green": color("38B06A"),
    "green_soft": color("D6F0E3"),
    "green_mid": color("2E9459"),
    "green_dark": color("217A4A"),
    "text": color("1A1A1A"),
    "text_secondary": color("4A4A4A"),
    "text_muted": color("7A7A7A"),
    "text_light": color("9E9E9E"),
    "surface": color("F5F5F5"),
    "border": color("E8E8E8"),
    "panel": color("EFEFEF"),
    "yellow": color("F5C518"),
    "yellow_light": color("FEF7D0"),
    "orange": color("F75801"),
    "gray_deco": color("B0B0B0"),
}


def add_shape(slide, shape_type, x, y, width, height, fill_color, line_color=None, transparency=0):
    shape = slide.shapes.add_shape(shape_type, x, y, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.fill.transparency = transparency
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    return shape


def add_rect(slide, x, y, width, height, fill_color, line_color=None, transparency=0):
    return add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        x,
        y,
        width,
        height,
        fill_color,
        line_color,
        transparency,
    )


def add_oval(slide, x, y, width, height, fill_color, transparency=0):
    return add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.OVAL,
        x,
        y,
        width,
        height,
        fill_color,
        None,
        transparency,
    )


def add_text(
    slide,
    text,
    x,
    y,
    width,
    height,
    font_size,
    font_name=DISPLAY_FONT,
    bold=True,
    font_color=None,
    align=PP_ALIGN.LEFT,
    valign=MSO_VERTICAL_ANCHOR.TOP,
    margin_left=0,
    margin_right=0,
    margin_top=0,
    margin_bottom=0,
):
    text_box = slide.shapes.add_textbox(x, y, width, height)
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = valign
    text_frame.margin_left = margin_left
    text_frame.margin_right = margin_right
    text_frame.margin_top = margin_top
    text_frame.margin_bottom = margin_bottom
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = font_color or COLORS["text"]
    return text_box


def add_header(slide, eyebrow, title, subtitle):
    add_text(
        slide,
        eyebrow.upper(),
        HEADER_X,
        EYEBROW_Y,
        HEADER_W,
        Inches(0.24),
        14,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text_light"],
    )
    add_text(
        slide,
        title,
        HEADER_X,
        TITLE_Y,
        HEADER_W,
        Inches(0.48),
        30,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text"],
    )
    add_text(
        slide,
        subtitle,
        HEADER_X,
        TITLE_Y + Inches(0.42),
        HEADER_W,
        Inches(0.24),
        15,
        font_name=DISPLAY_FONT,
        font_color=COLORS["green"],
    )


def set_background(slide, image_path):
    picture = slide.shapes.add_picture(image_path, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    element = picture._element
    element.getparent().remove(element)
    slide.shapes._spTree.insert(2, element)


def add_card(slide, x, y, width, height, accent_color=None, fill_color=None):
    fill = fill_color or COLORS["surface"]
    add_rect(slide, x, y, width, height, fill, COLORS["border"])
    if accent_color is not None:
        add_rect(slide, x, y, width, Pt(4), accent_color)


def copy_bundle_assets():
    os.makedirs(BUNDLE_BG_DIR, exist_ok=True)
    shutil.copy2(COVER_BG_SOURCE, COVER_BG)
    shutil.copy2(CONTENT_BG_SOURCE, CONTENT_BG)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COVER_BG)

    add_oval(slide, Inches(7.35), Inches(-0.25), Inches(2.55), Inches(2.55), COLORS["gray_deco"], 0.75)
    add_oval(slide, Inches(-0.30), Inches(4.35), Inches(1.40), Inches(1.40), COLORS["gray_deco"], 0.82)

    panel_x = Inches(1.00)
    panel_y = Inches(1.15)
    panel_w = Inches(7.20)
    panel_h = Inches(3.20)
    add_rect(slide, panel_x, panel_y, panel_w, panel_h, COLORS["white"], COLORS["white"])
    add_rect(slide, panel_x, panel_y, panel_w, Pt(6), COLORS["green"])

    add_text(
        slide,
        "DATA PRODUCT INTRO",
        panel_x + Inches(0.22),
        panel_y + Inches(0.18),
        Inches(3.60),
        Inches(0.20),
        14,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text_light"],
    )
    add_text(
        slide,
        "Yahoo Fantasy Baseball\nAnalyzer v2",
        panel_x + Inches(0.22),
        panel_y + Inches(0.52),
        Inches(5.80),
        Inches(1.25),
        40,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text"],
    )
    add_text(
        slide,
        "介紹一個支援 H2H 雙模式、球員排行、對位比較與交易模擬的單頁分析平台。",
        panel_x + Inches(0.22),
        panel_y + Inches(1.92),
        Inches(6.20),
        Inches(0.42),
        16,
        font_name=DISPLAY_FONT,
        font_color=COLORS["green"],
    )
    add_rect(slide, panel_x + Inches(0.22), panel_y + Inches(2.42), Inches(0.42), Pt(3), COLORS["green"])

    meta_items = [
        ("HTML", "Single file app"),
        ("MODES", "Categories / Points"),
        ("USE", "Draft, FA, Trade"),
    ]
    for index, (value, label) in enumerate(meta_items):
        meta_x = panel_x + Inches(0.22) + Inches(2.10) * index
        add_text(
            slide,
            value,
            meta_x,
            panel_y + Inches(2.62),
            Inches(1.60),
            Inches(0.26),
            16,
            font_name=MONO_FONT,
            font_color=COLORS["green"],
        )
        add_text(
            slide,
            label,
            meta_x,
            panel_y + Inches(2.90),
            Inches(1.85),
            Inches(0.26),
            14,
            font_name=DISPLAY_FONT,
            font_color=COLORS["text_secondary"],
        )


def slide_feature_grid(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CONTENT_BG)
    add_header(slide, "Feature Map", "這個 HTML 網頁提供哪些能力", "從資料匯入到交易決策，全部集中在同一個分析介面")

    items = [
        ("01", "雙模式", "H2H Categories 與 H2H Points 一鍵切換", COLORS["green_soft"]),
        ("02", "資料匯入", "拖拉 CSV / JSON、多檔合併、貼上解析", COLORS["surface"]),
        ("03", "打者排行", "搜尋、守位、Fantasy 類型、排序與 Tier", COLORS["surface"]),
        ("04", "投手排行", "SP / RP、風險權重、ERA / WHIP / K 類別", COLORS["surface"]),
        ("05", "KPI 篩選", "快速聚焦 elite、multitool、closer、ace", COLORS["yellow_light"]),
        ("06", "欄位客製", "欄位開關與 H2H 預設，表格可自由精簡", COLORS["surface"]),
        ("07", "觀察名單", "星號收藏與 localStorage 持久化", COLORS["surface"]),
        ("08", "對位分析", "雙球員卡片、雷達圖、洞察文字", COLORS["surface"]),
        ("09", "交易模擬", "Give / Receive 名單、價值差與風險差", COLORS["yellow_light"]),
        ("10", "回應式 UI", "深色 dashboard 視覺與桌機 / 行動版適配", COLORS["surface"]),
    ]

    card_gap_x = Inches(0.12)
    card_gap_y = Inches(0.18)
    card_w = (SAFE_W - card_gap_x * 4) / 5
    card_h = Inches(1.48)

    for index, (code, title, body, fill) in enumerate(items):
        row = index // 5
        col = index % 5
        card_x = SAFE_X + (card_w + card_gap_x) * col
        card_y = CONTENT_Y + (card_h + card_gap_y) * row
        accent = COLORS["green"] if index in {0, 4, 8} else COLORS["border"]
        add_card(slide, card_x, card_y, card_w, card_h, accent_color=accent, fill_color=fill)
        add_text(
            slide,
            code,
            card_x + Inches(0.10),
            card_y + Inches(0.12),
            Inches(0.50),
            Inches(0.30),
            22,
            font_name=MONO_FONT,
            font_color=COLORS["green"],
        )
        add_text(
            slide,
            title,
            card_x + Inches(0.10),
            card_y + Inches(0.44),
            card_w - Inches(0.20),
            Inches(0.26),
            14,
            font_name=DISPLAY_FONT,
            font_color=COLORS["text"],
        )
        add_text(
            slide,
            body,
            card_x + Inches(0.10),
            card_y + Inches(0.78),
            card_w - Inches(0.20),
            Inches(0.52),
            14,
            font_name=DISPLAY_FONT,
            font_color=COLORS["text_secondary"],
        )


def slide_runtime_panels(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CONTENT_BG)
    add_header(slide, "Runtime Model", "核心運作原理", "資料流、狀態與視圖重新整理的關係，全部寫在單一 HTML 內")

    panel_ratios = [3.2, 1.1, 1.1, 2.9]
    panel_gap = Inches(0.12)
    available_width = SAFE_W - panel_gap * 3
    ratio_sum = sum(panel_ratios)
    panel_widths = [available_width * ratio / ratio_sum for ratio in panel_ratios]
    panel_x = SAFE_X
    panel_y = CONTENT_Y
    panel_h = Inches(3.65)

    panels = [
        {
            "number": "01",
            "title": "Input Layer",
            "body": "支援預設資料、CSV / JSON 上傳、文字貼上與本地 Yahoo 匯出自動載入。\nparseCsv / parseJson / applyParsedData 會把資料整理成 batters 與 pitchers 兩個 pool。",
            "fill": COLORS["surface"],
            "accent": COLORS["green"],
        },
        {
            "number": "02",
            "title": "Modes",
            "body": "CATS\nPOINTS",
            "fill": COLORS["green_soft"],
            "accent": COLORS["green"],
        },
        {
            "number": "04",
            "title": "Tabs",
            "body": "BAT\nPIT\nCMP\nTRD",
            "fill": COLORS["yellow_light"],
            "accent": COLORS["yellow"],
        },
        {
            "number": "LOOP",
            "title": "Decision Loop",
            "body": "computeAllScores() 重新計算 valueScore / riskScore，\nrebuildUi() 再依序 renderBatters、renderPitchers、renderCompare、calcTrade。\n模式、排序、欄位、KPI、觀察名單都會回到同一個狀態中心。",
            "fill": COLORS["surface"],
            "accent": COLORS["green_mid"],
        },
    ]

    for index, panel in enumerate(panels):
        width = panel_widths[index]
        add_card(slide, panel_x, panel_y, width, panel_h, accent_color=panel["accent"], fill_color=panel["fill"])
        add_text(
            slide,
            panel["number"],
            panel_x + Inches(0.16),
            panel_y + Inches(0.18),
            width - Inches(0.32),
            Inches(0.38),
            28,
            font_name=MONO_FONT,
            font_color=panel["accent"],
        )
        add_text(
            slide,
            panel["title"],
            panel_x + Inches(0.16),
            panel_y + Inches(0.62),
            width - Inches(0.32),
            Inches(0.30),
            16,
            font_name=DISPLAY_FONT,
            font_color=COLORS["text"],
        )
        add_text(
            slide,
            panel["body"],
            panel_x + Inches(0.16),
            panel_y + Inches(1.02),
            width - Inches(0.32),
            Inches(2.30),
            14,
            font_name=DISPLAY_FONT if index in {0, 3} else MONO_FONT,
            font_color=COLORS["text_secondary"],
        )
        panel_x += width + panel_gap


def add_rank_row(slide, x, y, width, rank, title, label, value, accent_color):
    row_h = Inches(0.58)
    add_card(slide, x, y, width, row_h, accent_color=COLORS["border"], fill_color=COLORS["surface"])
    add_rect(slide, x + Inches(0.10), y + Inches(0.10), Inches(0.36), Inches(0.36), accent_color)
    add_text(
        slide,
        rank,
        x + Inches(0.10),
        y + Inches(0.12),
        Inches(0.36),
        Inches(0.26),
        14,
        font_name=MONO_FONT,
        font_color=COLORS["white"],
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        title,
        x + Inches(0.56),
        y + Inches(0.08),
        Inches(2.10),
        Inches(0.20),
        14,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text"],
    )
    add_text(
        slide,
        label,
        x + Inches(0.56),
        y + Inches(0.30),
        Inches(1.90),
        Inches(0.18),
        14,
        font_name=MONO_FONT,
        font_color=COLORS["text_light"],
    )
    add_text(
        slide,
        f"{value}",
        x + width - Inches(0.82),
        y + Inches(0.08),
        Inches(0.72),
        Inches(0.20),
        14,
        font_name=MONO_FONT,
        font_color=accent_color,
        align=PP_ALIGN.RIGHT,
    )
    add_rect(slide, x + width - Inches(0.78), y + Inches(0.34), Inches(0.62), Pt(4), COLORS["border"])
    add_rect(
        slide,
        x + width - Inches(0.78),
        y + Inches(0.34),
        Inches(0.62) * value / 100,
        Pt(4),
        accent_color,
    )


def slide_highlight_ranking(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CONTENT_BG)
    add_header(slide, "Top Highlights", "最值得 demo 的 10 個頁面亮點", "依實際互動價值排序，方便口頭展示時抓節奏")

    rows = [
        ("01", "Mode Switch", "ENTRY / RULESET", 96, COLORS["green"]),
        ("02", "Import Zone", "DATA / CSV / JSON", 94, COLORS["green"]),
        ("03", "Batter Filters", "BAT / SEARCH / KPI", 92, COLORS["green_mid"]),
        ("04", "Pitcher Risk Tuning", "PIT / WEIGHTS", 90, COLORS["green_mid"]),
        ("05", "Tiered Table", "VAL / RISK / TIER", 88, COLORS["green_dark"]),
        ("06", "Column Presets", "TABLE / H2H", 86, COLORS["yellow"]),
        ("07", "Watchlist Star", "LOCAL STORAGE", 84, COLORS["yellow"]),
        ("08", "Radar Compare", "CMP / CANVAS", 82, COLORS["green"]),
        ("09", "Trade Verdict", "TRD / WIN LOSE", 80, COLORS["orange"]),
        ("10", "Auto Load CSV", "BOOT / FILE IO", 78, COLORS["orange"]),
    ]

    column_gap = Inches(0.18)
    column_w = (SAFE_W - column_gap) / 2
    row_gap = Inches(0.12)
    row_y = CONTENT_Y

    for index, row in enumerate(rows):
        column = 0 if index < 5 else 1
        offset = index if index < 5 else index - 5
        x = SAFE_X + column * (column_w + column_gap)
        y = row_y + offset * (Inches(0.58) + row_gap)
        add_rank_row(slide, x, y, column_w, *row)


def slide_tab_cards(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CONTENT_BG)
    add_header(slide, "Primary Tabs", "四個主頁籤，對應四種決策任務", "每張卡片代表一個核心使用情境")

    card_gap = Inches(0.14)
    card_w = (SAFE_W - card_gap * 3) / 4
    card_h = Inches(3.55)
    cards = [
        ("BAT", "打者排行", ["搜尋球員或球隊", "依守位與 Fantasy 類型過濾", "查看 AVG / OBP / OPS / SB", "用 KPI 快速找跑打兼具與 elite", "支援欄位顯示切換"], "找打者補強與撿 FA 的主畫面", COLORS["green"]),
        ("PIT", "投手排行", ["SP / RP 分開篩選", "可調 SP / RP 風險權重", "查看 ERA / WHIP / K/9 / QS / SVH", "用 closer / ace KPI 快速聚焦", "評分模式隨玩法同步切換"], "處理投手配置與風險控制", COLORS["green_mid"]),
        ("CMP", "對位分析", ["球員 A / B 可各自選打者或投手", "搜尋即時縮小候選名單", "卡片摘要顯示關鍵資料", "Canvas 雷達圖與長條圖比較", "insight 區給出文字洞察"], "用來處理二選一與 waiver 決策", COLORS["yellow"]),
        ("TRD", "交易模擬", ["Give / Receive 兩側清單", "搜尋後直接加入名單", "即時重算總價值與平均風險", "用 verdict box 顯示 Win / Lose / Even", "橫條圖呈現雙方落差"], "用於談判前的交易估值", COLORS["orange"]),
    ]

    for index, (code, title, lines, footer, accent) in enumerate(cards):
        card_x = SAFE_X + index * (card_w + card_gap)
        card_y = CONTENT_Y
        add_card(slide, card_x, card_y, card_w, card_h, accent_color=accent, fill_color=COLORS["surface"])
        add_text(
            slide,
            code,
            card_x + Inches(0.12),
            card_y + Inches(0.18),
            card_w - Inches(0.24),
            Inches(0.32),
            28,
            font_name=MONO_FONT,
            font_color=accent,
        )
        add_text(
            slide,
            title,
            card_x + Inches(0.12),
            card_y + Inches(0.56),
            card_w - Inches(0.24),
            Inches(0.24),
            15,
            font_name=DISPLAY_FONT,
            font_color=COLORS["text"],
        )
        for line_index, line_text in enumerate(lines):
            add_text(
                slide,
                line_text,
                card_x + Inches(0.12),
                card_y + Inches(0.98) + Inches(0.38) * line_index,
                card_w - Inches(0.24),
                Inches(0.26),
                14,
                font_name=DISPLAY_FONT,
                font_color=COLORS["text_secondary"],
            )
        add_rect(slide, card_x + Inches(0.12), card_y + Inches(3.00), card_w - Inches(0.24), Pt(1), COLORS["border"])
        add_text(
            slide,
            footer,
            card_x + Inches(0.12),
            card_y + Inches(3.10),
            card_w - Inches(0.24),
            Inches(0.34),
            14,
            font_name=DISPLAY_FONT,
            font_color=COLORS["text_light"],
        )


def slide_visual_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CONTENT_BG)
    add_header(slide, "Visual + State", "頁面風格與前端架構", "左邊講設計語言，右邊把 HTML 內的狀態與 render 關係畫出來")

    left_w = Inches(3.05)
    right_w = SAFE_W - left_w - Inches(0.24)
    top_h = Inches(1.55)
    bottom_h = Inches(1.80)

    add_card(slide, SAFE_X, CONTENT_Y, left_w, top_h, accent_color=COLORS["green"], fill_color=COLORS["surface"])
    add_text(slide, "視覺設計語言", SAFE_X + Inches(0.16), CONTENT_Y + Inches(0.16), left_w - Inches(0.32), Inches(0.24), 15, font_name=DISPLAY_FONT, font_color=COLORS["text"])
    add_text(
        slide,
        "深色 dashboard 背景 + 亮色發光重點\nhero、tab、panel、table 四層結構分明\n以 CSS 變數維持色彩與間距一致\n桌機與行動版都保留主要分析流程",
        SAFE_X + Inches(0.16),
        CONTENT_Y + Inches(0.48),
        left_w - Inches(0.32),
        Inches(0.90),
        14,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text_secondary"],
    )

    lower_y = CONTENT_Y + top_h + Inches(0.14)
    add_card(slide, SAFE_X, lower_y, left_w, bottom_h, accent_color=COLORS["yellow"], fill_color=COLORS["surface"])
    add_text(slide, "資料模型與互動狀態", SAFE_X + Inches(0.16), lower_y + Inches(0.16), left_w - Inches(0.32), Inches(0.24), 15, font_name=DISPLAY_FONT, font_color=COLORS["text"])
    add_text(
        slide,
        "batters / pitchers 作為主資料池\nbState / pState 管控搜尋、排序、KPI 篩選\ntradeState 管理交易雙方名單\nwatchlist 與 localStorage 讓標記結果可持久保存",
        SAFE_X + Inches(0.16),
        lower_y + Inches(0.48),
        left_w - Inches(0.32),
        Inches(1.05),
        14,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text_secondary"],
    )

    right_x = SAFE_X + left_w + Inches(0.24)
    add_card(slide, right_x, CONTENT_Y, right_w, Inches(3.49), accent_color=COLORS["green_mid"], fill_color=COLORS["surface"])
    add_text(slide, "APP FLOW / ONE FILE", right_x + Inches(0.18), CONTENT_Y + Inches(0.18), right_w - Inches(0.36), Inches(0.22), 15, font_name=MONO_FONT, font_color=COLORS["green_mid"])

    top_box_y = CONTENT_Y + Inches(0.52)
    box_gap = Inches(0.10)
    mini_w = (right_w - Inches(0.36) - box_gap * 2) / 3
    mini_labels = [
        ("INPUT", "fileInput\npasteArea\nautoLoadLocalExports", COLORS["green_soft"]),
        ("STATE", "bState\npState\ntradeState\nwatchlist", COLORS["yellow_light"]),
        ("SCORE", "computeBatterScores\ncomputePitcherScores\ncomputeAllScores", COLORS["green_soft"]),
    ]
    for index, (title, body, fill) in enumerate(mini_labels):
        mini_x = right_x + Inches(0.18) + index * (mini_w + box_gap)
        add_card(slide, mini_x, top_box_y, mini_w, Inches(0.92), accent_color=COLORS["border"], fill_color=fill)
        add_text(slide, title, mini_x + Inches(0.10), top_box_y + Inches(0.10), mini_w - Inches(0.20), Inches(0.18), 14, font_name=MONO_FONT, font_color=COLORS["text"])
        add_text(slide, body, mini_x + Inches(0.10), top_box_y + Inches(0.34), mini_w - Inches(0.20), Inches(0.48), 14, font_name=MONO_FONT, font_color=COLORS["text_secondary"])

    core_y = top_box_y + Inches(1.10)
    add_card(slide, right_x + Inches(0.18), core_y, right_w - Inches(0.36), Inches(0.86), accent_color=COLORS["green"], fill_color=COLORS["panel"])
    add_text(slide, "rebuildUi()  ->  renderBatters() / renderPitchers() / renderCompare() / calcTrade()", right_x + Inches(0.30), core_y + Inches(0.28), right_w - Inches(0.60), Inches(0.24), 14, font_name=MONO_FONT, font_color=COLORS["text"])

    bottom_y = core_y + Inches(1.08)
    bottom_w = (right_w - Inches(0.36) - box_gap * 2) / 3
    bottom_items = [
        ("BATTERS VIEW", "table + KPI filters + watch star", COLORS["surface"]),
        ("PITCHERS VIEW", "risk tuning + table presets", COLORS["surface"]),
        ("DECISION VIEW", "compare canvas + trade verdict", COLORS["surface"]),
    ]
    for index, (title, body, fill) in enumerate(bottom_items):
        box_x = right_x + Inches(0.18) + index * (bottom_w + box_gap)
        add_card(slide, box_x, bottom_y, bottom_w, Inches(0.92), accent_color=COLORS["border"], fill_color=fill)
        add_text(slide, title, box_x + Inches(0.10), bottom_y + Inches(0.10), bottom_w - Inches(0.20), Inches(0.18), 14, font_name=MONO_FONT, font_color=COLORS["text"])
        add_text(slide, body, box_x + Inches(0.10), bottom_y + Inches(0.36), bottom_w - Inches(0.20), Inches(0.34), 14, font_name=DISPLAY_FONT, font_color=COLORS["text_secondary"])


def slide_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CONTENT_BG)
    add_oval(slide, Inches(8.05), Inches(4.05), Inches(2.10), Inches(2.10), COLORS["gray_deco"], 0.82)

    add_text(slide, "TAKEAWAY", HEADER_X, Inches(0.58), HEADER_W, Inches(0.22), 14, font_name=DISPLAY_FONT, font_color=COLORS["text_light"])
    add_text(slide, "一個單檔 HTML，完成 Fantasy Baseball 的完整決策鏈", HEADER_X, Inches(0.92), HEADER_W, Inches(0.66), 32, font_name=DISPLAY_FONT, font_color=COLORS["text"])
    add_text(
        slide,
        "它把資料匯入、雙模式評分、排行榜、球員比較與交易模擬整合在同一個介面，\n適合拿來展示資料產品思維、前端狀態管理，以及夢幻聯盟決策工具的產品化方式。",
        HEADER_X,
        Inches(1.78),
        HEADER_W,
        Inches(0.70),
        16,
        font_name=DISPLAY_FONT,
        font_color=COLORS["text_secondary"],
    )
    add_rect(slide, HEADER_X, Inches(2.62), Inches(0.42), Pt(3), COLORS["green"])

    stats = [
        ("2", "GAME MODES"),
        ("4", "PRIMARY TABS"),
        ("2", "PLAYER POOLS"),
        ("1", "HTML FILE"),
    ]
    stat_w = Inches(1.65)
    stat_gap = Inches(0.22)
    stat_y = Inches(3.65)
    for index, (value, label) in enumerate(stats):
        stat_x = HEADER_X + index * (stat_w + stat_gap)
        add_text(slide, value, stat_x, stat_y, stat_w, Inches(0.38), 40, font_name=MONO_FONT, font_color=COLORS["green"])
        add_text(slide, label, stat_x, stat_y + Inches(0.48), stat_w, Inches(0.20), 14, font_name=DISPLAY_FONT, font_color=COLORS["text_light"])


def build_manifest():
    manifest = {
        "deck": DECK_NAME,
        "template": "fetch_template_with_backgrounds_editorial.yaml",
        "generated": "2026-05-25",
        "slides": 7,
        "sources": {
            "html": os.path.relpath(HTML_SOURCE, BASE_DIR),
            "yaml": os.path.relpath(YAML_SOURCE, BASE_DIR),
        },
        "resources": [
            {
                "type": "background",
                "role": "cover",
                "source": os.path.relpath(COVER_BG_SOURCE, BASE_DIR),
                "bundled_as": "assets/backgrounds/cover-background.png",
            },
            {
                "type": "background",
                "role": "content",
                "source": os.path.relpath(CONTENT_BG_SOURCE, BASE_DIR),
                "bundled_as": "assets/backgrounds/content-background.png",
            },
        ],
    }
    with open(os.path.join(BUNDLE_ASSET_DIR, "resource-manifest.json"), "w", encoding="utf-8") as handle:
        json.dump(manifest, handle, indent=2, ensure_ascii=False)


def main():
    copy_bundle_assets()

    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    slide_cover(presentation)
    slide_feature_grid(presentation)
    slide_runtime_panels(presentation)
    slide_highlight_ranking(presentation)
    slide_tab_cards(presentation)
    slide_visual_architecture(presentation)
    slide_closing(presentation)

    presentation.save(OUTPUT_PPTX)
    build_manifest()

    print(f"[+] Bundle: {BUNDLE_DIR}")
    print(f"[+] PPTX: {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()